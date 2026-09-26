import os
import argparse
import numpy as np
import torch
from tqdm import tqdm

from data.dataset import readIndex, dataReadPip, loadedDataset
from model.deepcrack import DeepCrack
from config import Config as cfg

import os
from datetime import datetime
from openpyxl import Workbook, load_workbook


def append_metrics_to_excel(
    xlsx_path: str,
    model_name: str,
    metrics: dict,
    ckpt_path: str,
    val_list_path: str,
    threshold: float,
    sheet_name: str = "FinalPerformance"
):
    xlsx_path = os.path.abspath(xlsx_path)
    os.makedirs(os.path.dirname(xlsx_path), exist_ok=True)

    if os.path.exists(xlsx_path):
        wb = load_workbook(xlsx_path)
    else:
        wb = Workbook()
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
            ws0 = wb["Sheet"]
            wb.remove(ws0)

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)

    if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
        headers = [
            "Timestamp",
            "Model",
            "Precision",
            "Recall",
            "F1",
            "IoU",
            "TP",
            "FP",
            "FN",
            "Threshold",
            "ValList",
            "Checkpoint"
        ]
        ws.append(headers)

    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    row = [
        ts,
        model_name,
        float(metrics["Precision"]),
        float(metrics["Recall"]),
        float(metrics["F1"]),
        float(metrics["IoU"]),
        int(metrics["TP"]),
        int(metrics["FP"]),
        int(metrics["FN"]),
        float(threshold),
        os.path.abspath(val_list_path),
        os.path.abspath(ckpt_path),
    ]
    ws.append(row)

    wb.save(xlsx_path)
    print(f"Metrics appended to Excel: {xlsx_path} (sheet: {sheet_name})")


def binarize_gt(gt: torch.Tensor) -> torch.Tensor:
    return (gt > 0).to(torch.uint8)


def binarize_pred(logits: torch.Tensor, th: float) -> torch.Tensor:
    prob = torch.sigmoid(logits)
    return (prob[:, 0, :, :] > th).to(torch.uint8)


@torch.no_grad()
def eval_pixel_metrics(model, loader, device, th=0.5, save_pred_dir=None, names_list=None):
    TP = 0
    FP = 0
    FN = 0

    if save_pred_dir is not None:
        os.makedirs(save_pred_dir, exist_ok=True)

    idx_global = 0

    for batch in tqdm(loader, desc="Evaluating (pixel-level micro)"):
        img, gt = batch  # img: (B,3,H,W), gt: (B,H,W)
        img = img.to(device)
        gt = gt.to(device)

        outputs = model(img)
        pred_logits = outputs[0]  # (B,1,H,W)

        pred_bin = binarize_pred(pred_logits, th=th)  # (B,H,W) 0/1
        gt_bin = binarize_gt(gt)                      # (B,H,W) 0/1

        # TP: pred=1, gt=1
        TP += torch.logical_and(pred_bin == 1, gt_bin == 1).sum().item()
        # FP: pred=1, gt=0
        FP += torch.logical_and(pred_bin == 1, gt_bin == 0).sum().item()
        # FN: pred=0, gt=1
        FN += torch.logical_and(pred_bin == 0, gt_bin == 1).sum().item()

        if save_pred_dir is not None:
            pred_to_save = (pred_bin.cpu().numpy() * 255).astype(np.uint8)  # (B,H,W)
            bsz = pred_to_save.shape[0]
            for i in range(bsz):
                if names_list is not None and idx_global < len(names_list):
                    img_path = names_list[idx_global][0]
                    base = os.path.splitext(os.path.basename(img_path))[0]
                    out_name = f"{base}_pred.png"
                else:
                    out_name = f"{idx_global:06d}_pred.png"
                out_path = os.path.join(save_pred_dir, out_name)
                try:
                    import imageio.v2 as imageio
                    imageio.imwrite(out_path, pred_to_save[i])
                except Exception:
                    from PIL import Image
                    Image.fromarray(pred_to_save[i]).save(out_path)
                idx_global += 1

    eps = 1e-12
    precision = TP / (TP + FP + eps)
    recall = TP / (TP + FN + eps)
    iou = TP / (TP + FP + FN + eps)
    f1 = (2 * TP) / (2 * TP + FP + FN + eps)

    return {
        "TP": TP, "FP": FP, "FN": FN,
        "Precision": precision,
        "Recall": recall,
        "IoU": iou,
        "F1": f1
    }


def main():
    DEFAULT_CKPT = "checkpoints/MyDeepCrack2/checkpoints/MyDeepCrack2_MyDeepCrack2_epoch(149)_0000232_2025-07-15-18-05-54.pth"
    DEFAULT_LIST = "data/path_local_txt/val.txt"

    parser = argparse.ArgumentParser()
    parser.add_argument("--ckpt", type=str, default=DEFAULT_CKPT, help="path to .pth model weights")
    parser.add_argument("--list", type=str, default=DEFAULT_LIST, help="val list txt")
    parser.add_argument("--th", type=float, default=None, help="sigmoid threshold (default: cfg.acc_sigmoid_th)")
    parser.add_argument("--batch", type=int, default=1, help="eval batch size")
    parser.add_argument("--num_workers", type=int, default=None, help="dataloader workers (default: cfg.num_workers)")
    parser.add_argument("--save_pred", type=str, default=None, help="optional dir to save binary predictions")
    args = parser.parse_args()

    val_list_path = args.list if args.list is not None else cfg.val_data_path
    th = args.th if args.th is not None else getattr(cfg, "acc_sigmoid_th", 0.5)
    num_workers = args.num_workers if args.num_workers is not None else cfg.num_workers

    if not os.path.isfile(args.ckpt):
        raise FileNotFoundError(f"ckpt not found: {args.ckpt}")
    if not os.path.isfile(val_list_path):
        raise FileNotFoundError(f"val list txt not found: {val_list_path}")

    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")

    pip = dataReadPip(transforms=None)
    val_list = readIndex(val_list_path, shuffle=False)
    val_dataset = loadedDataset(val_list, preprocess=pip)

    val_loader = torch.utils.data.DataLoader(
        val_dataset,
        batch_size=args.batch,
        shuffle=False,
        num_workers=num_workers,
        drop_last=False
    )

    # ---- model ----
    model = DeepCrack().to(device)

    state = torch.load(args.ckpt, map_location=device)
    if isinstance(state, dict) and "state_dict" in state:
        state = state["state_dict"]

    new_state = {}
    for k, v in state.items():
        if k.startswith("module."):
            new_state[k[len("module."):]] = v
        else:
            new_state[k] = v
    model.load_state_dict(new_state, strict=True)
    model.eval()

    metrics = eval_pixel_metrics(
        model=model,
        loader=val_loader,
        device=device,
        th=th,
        save_pred_dir=args.save_pred,
        names_list=val_list
    )

    print("\n===== DeepCrack Final Pixel-level Metrics (micro, val set) =====")
    print(f"Val list: {val_list_path}")
    print(f"Checkpoint: {args.ckpt}")
    print(f"Threshold: {th}")
    print(f"TP={metrics['TP']}, FP={metrics['FP']}, FN={metrics['FN']}")
    print(f"Precision = {metrics['Precision']:.6f}")
    print(f"Recall    = {metrics['Recall']:.6f}")
    print(f"F1-score  = {metrics['F1']:.6f}")
    print(f"IoU       = {metrics['IoU']:.6f}")
    print("==============================================================\n")
    EXCEL_PATH = os.path.abspath("final_performance_comparison.xlsx")

    append_metrics_to_excel(
        xlsx_path=EXCEL_PATH,
        model_name="DeepCrack",
        metrics=metrics,
        ckpt_path=args.ckpt,
        val_list_path=val_list_path,
        threshold=th,
        sheet_name="FinalPerformance"
    )

if __name__ == "__main__":
    main()
